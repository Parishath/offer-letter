import pandas as pd
from docx import Document
from docx2pdf import convert
import os

# Paths
# Path to your Excel file
excel_path = "/offer-letter/doc/offer-ParvaM.xlsx"
# Path to your Word template
template_path = "/offer-letter/doc/Sahana H.docx"
output_dir = "Generated_Documents"  # Output directory for generated documents

# Read Excel data
data = pd.read_excel(excel_path)

# Create output directory if it doesn't exist
os.makedirs(output_dir, exist_ok=True)

# Generate personalized documents
for index, row in data.iterrows():
    name = row["Name"]  # Get the name from the Excel column

    # Open the Word template
    doc = Document(template_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        if "Sahana H" in paragraph.text:
            paragraph.text = paragraph.text.replace("Sahana H", name)

    # Save the document as .docx
    docx_filename = f"{output_dir}/{name}_Offer_Letter.docx"
    doc.save(docx_filename)

    # Convert to PDF
    pdf_filename = f"{output_dir}/{name}_Offer_Letter.pdf"
    convert(docx_filename, pdf_filename)

    print(f"Generated: {pdf_filename}")

print("All documents generated successfully!")
